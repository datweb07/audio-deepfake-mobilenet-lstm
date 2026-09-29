from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\audio-deepfake-mobilenet-lstm")
SOURCE = ROOT / "papers" / "PaperID 804 final.docx"
OUTPUT = ROOT / "BIT" / "BIT_final.docx"

FONT = "Times New Roman"
FONT_SIZE = Pt(13)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def set_text(paragraph, text: str) -> None:
    ppr = deepcopy(paragraph._p.pPr) if paragraph._p.pPr is not None else None
    for child in list(paragraph._p):
        paragraph._p.remove(child)
    if ppr is not None:
        paragraph._p.append(ppr)
    paragraph.add_run(text)


def find_exact(doc: Document, text: str):
    matches = [p for p in doc.paragraphs if p.text.strip() == text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph {text!r}, found {len(matches)}")
    return matches[0]


def add_before(doc: Document, target, text: str = "", style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    if text:
        paragraph.add_run(text)
    target._p.addprevious(paragraph._p)
    return paragraph


def add_page_break_before(doc: Document, target):
    paragraph = add_before(doc, target)
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    return paragraph


def add_field_run(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = ""
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, placeholder, end):
        run._r.append(element)


def set_run_font_xml(run_element) -> None:
    rpr = run_element.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run_element.insert(0, rpr)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT)
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        attr = qn(f"w:{key}")
        rfonts.attrib.pop(attr, None)

    size = rpr.find(qn("w:sz"))
    if size is None:
        size = OxmlElement("w:sz")
        rpr.append(size)
    size.set(qn("w:val"), "26")
    size_cs = rpr.find(qn("w:szCs"))
    if size_cs is None:
        size_cs = OxmlElement("w:szCs")
        rpr.append(size_cs)
    size_cs.set(qn("w:val"), "26")

    underline = rpr.find(qn("w:u"))
    if underline is not None:
        rpr.remove(underline)
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rpr.append(color)
    color.set(qn("w:val"), "000000")
    color.attrib.pop(qn("w:themeColor"), None)


def normalize_typography(doc: Document) -> None:
    # Style inheritance.
    for style in doc.styles:
        if style.type not in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            continue
        style.font.name = FONT
        style.font.size = FONT_SIZE
        style.font.color.rgb = RGBColor(0, 0, 0)
        if style.element.rPr is not None:
            rfonts = style.element.rPr.get_or_add_rFonts()
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                rfonts.set(qn(f"w:{key}"), FONT)

    roots = [doc.element]
    for section in doc.sections:
        roots.extend((section.header._element, section.footer._element))
    seen = set()
    for root in roots:
        if id(root) in seen:
            continue
        seen.add(id(root))
        for run in root.xpath(".//w:r"):
            set_run_font_xml(run)

    # Paragraph geometry, including tables and captions.
    for paragraph in doc.paragraphs:
        fmt = paragraph.paragraph_format
        fmt.line_spacing = 1.3
        fmt.space_after = Pt(4)
        fmt.space_before = Pt(0)
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            fmt.keep_with_next = True
            fmt.space_before = Pt(8)
            fmt.space_after = Pt(4)
            for run in paragraph.runs:
                run.bold = True
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.3
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.space_before = Pt(0)


def configure_pages(doc: Document) -> None:
    for index, section in enumerate(doc.sections):
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(0.7)
        section.footer_distance = Cm(0.7)
        section.different_first_page_header_footer = False

        section.header.is_linked_to_previous = False
        header = section.header
        for child in list(header._element):
            header._element.remove(child)
        paragraph = header.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field_run(paragraph, "PAGE")

        section.footer.is_linked_to_previous = False
        footer = section.footer
        for child in list(footer._element):
            footer._element.remove(child)
        footer.add_paragraph()

    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def style_frontmatter(paragraph, *, bold=False, center=False, upper=False):
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.bold = bold
        if upper:
            run.text = run.text.upper()


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    title = doc.paragraphs[0].text.strip()
    abstract = find_exact(doc, next(p.text for p in doc.paragraphs if p.style.name == "abstract")).text
    abstract = abstract.removeprefix("Abstract.").strip()
    keywords = find_exact(doc, next(p.text for p in doc.paragraphs if p.style.name == "keywords")).text

    # Remove the conference title/author/address/abstract block from the copied manuscript.
    for paragraph in list(doc.paragraphs[:6]):
        remove_paragraph(paragraph)

    # Remove acknowledgement as required by the anonymous BIT report rules.
    ack_heading = find_exact(doc, "Acknowledgement")
    paragraphs = doc.paragraphs
    ack_index = next(i for i, p in enumerate(paragraphs) if p._p is ack_heading._p)
    ack_body = paragraphs[ack_index + 1]
    remove_paragraph(ack_body)
    remove_paragraph(ack_heading)

    # Explicit, stable section numbering required by the competition.
    heading_map = {
        "Introduction": "1. Introduction",
        "Related Work": "2. Related Work",
        "Deepfake Voice Detection and Benchmark Protocols": "2.1. Deepfake Voice Detection and Benchmark Protocols",
        "Robustness Replay and Unseen Data Generalization": "2.2. Robustness, Replay, and Unseen-Data Generalization",
        "Computational Efficiency and Deployment Oriented Evaluation": "2.3. Computational Efficiency and Deployment-Oriented Evaluation",
        "Methodology": "3. Methodology",
        "LAVA Framework and Experimental Protocol": "3.1. LAVA System Design and Experimental Protocol",
        "Dataset Integrity and Canonical Split": "3.2. Dataset Integrity and Canonical Split",
        "Standardized Audio Preprocessing": "3.3. Standardized Audio Preprocessing",
        "Evaluated Detector Architectures": "3.4. Evaluated Detector Architectures",
        "Training Artifact Provenance and Score Semantics": "3.5. Training, Artifact Provenance, and Score Semantics",
        "Clean and Diagnostic Stress Evaluation": "3.6. Clean and Diagnostic Stress Evaluation",
        "Computational Efficiency Evaluation": "3.7. Computational Efficiency Evaluation",
        "Results and Discussion": "4. Results and Discussion",
        "RQ1 Clean Detection Performance": "4.1. RQ1: Clean Detection Performance",
        "RQ2 Diagnostic Performance under Executed Distortions": "4.2. RQ2: Diagnostic Performance under Executed Distortions",
        "Computational Efficiency": "4.3. Computational Efficiency",
        "RQ3 Corrected Pareto Trade-offs": "4.4. RQ3: Corrected Pareto Trade-offs",
        "Conclusion": "5. Conclusions and Recommendations",
        "Practical and Societal Implications": "5.1. Practical and Societal Implications",
        "Limitations": "5.2. Limitations",
        "Concluding Remarks": "5.3. Conclusions and Future Work",
    }
    for old, new in heading_map.items():
        set_text(find_exact(doc, old), new)

    # Remove an empty Heading 2 paragraph inherited from the conference layout;
    # otherwise it would create a blank entry in navigation/TOC views.
    for paragraph in list(doc.paragraphs):
        if paragraph.style and paragraph.style.name == "heading2" and not paragraph.text.strip():
            remove_paragraph(paragraph)

    broken_caption = find_exact(
        doc,
        "Table 4. Fixed 100-recording diagnostic stress results. ΔF1 is clean minus stressed perfor-mance; replay is simulated.",
    )
    set_text(
        broken_caption,
        "Table 4. Fixed 100-recording diagnostic stress results. ΔF1 is clean minus stressed performance; replay is simulated.",
    )

    # Replace malformed symbolic bullets by an ordinary numbered contribution list.
    contribution_texts = [
        "1) A checksum-group-disjoint, artifact-traceable clean evaluation contract centered on ROC-AUC/EER, calibrated score semantics, and explicit provenance.",
        "2) Matched fixed-subset diagnostics for multiple noise levels, codec settings, and simulated replay, with aggregate failure-pattern analysis.",
        "3) Single-thread desktop-CPU timing and exploratory three-objective Pareto analysis for scenario-specific artifact selection.",
    ]
    malformed = [p for p in doc.paragraphs if p.text.strip().startswith("\u2022")]
    if len(malformed) != 3:
        raise RuntimeError(f"Expected three contribution bullets, found {len(malformed)}")
    for paragraph, text in zip(malformed, contribution_texts):
        set_text(paragraph, text)

    contribution_lead = find_exact(doc, "The study makes three evidence-backed contributions:")
    set_text(contribution_lead, "1.2. Contributions")
    contribution_lead.style = doc.styles["heading2"]

    rq1 = find_exact(doc, "RQ1: How do deployed artifacts compare on clean AUC and EER, and how should threshold-dependent metrics be interpreted under heterogeneous training and calibration?")
    scope_heading = rq1.insert_paragraph_before("1.1. Research Objectives, Object, Method, and Scope", style="heading2")
    scope_text = (
        "The objective is to design and empirically evaluate an artifact-traceable framework for comparing heterogeneous deepfake-voice detectors under shared integrity, score, diagnostic-stress, and runtime contracts. "
        "The research objects are six deployed detector artifacts and their end-to-end inference pipelines. The method combines SHA-256 data-integrity auditing, canonical clean evaluation, matched diagnostic stress tests, measured CPU efficiency, bootstrap and paired statistical analysis, and multi-objective Pareto analysis. "
        "The scope comprises all 2,737 canonical test recordings for clean evaluation, a fixed 100-recording subset for diagnostic stress, and one Windows desktop CPU; it does not include an unseen corpus, physical replay, causal streaming, representative edge hardware, or multi-seed retraining. The research questions are:"
    )
    rq1.insert_paragraph_before(scope_text, style="p1a")

    framework = find_exact(doc, "LAVA separates canonical data and integrity, registered detectors, conditions, measurements, and analyses. The registry records each artifact's input, duration, threshold, and provenance; adapters preserve native computation while exposing common scores and resource measures. Figure 1 summarizes the six-detector contract, and Figure 2 shows the integrity-controlled data flow.")
    set_text(
        framework,
        framework.text
        + " The production implementation contains an artifact-gated detector registry, preprocessing and framework adapters, clean/stress/efficiency runners, aggregate reporting modules, and a responsive Streamlit dashboard; the interface loads validated artifacts and never trains models.",
    )

    # Front matter, inserted before Section 1.
    intro = find_exact(doc, "1. Introduction")
    cover_lines = [
        ("UNIVERSITY OF ECONOMICS HO CHI MINH CITY", True, True),
        ("SCHOOL OF BUSINESS INFORMATION TECHNOLOGY", True, True),
        ("", False, True),
        ("BIT GENESIS RESEARCH AWARDS 2026", True, True),
        ("FULL RESEARCH REPORT", True, True),
        ("", False, True),
        (title, True, True),
        ("", False, True),
        ("Field: Information Technology and Engineering", False, True),
        ("Ho Chi Minh City, 2026", False, True),
    ]
    for text, bold, center in cover_lines:
        p = add_before(doc, intro, text)
        style_frontmatter(p, bold=bold, center=center)
        if text == title:
            p.paragraph_format.space_before = Pt(36)
            p.paragraph_format.space_after = Pt(36)
    add_page_break_before(doc, intro)

    summary_heading = add_before(doc, intro, "RESEARCH SUMMARY", style="heading1")
    summary_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_before(doc, intro, abstract, style="Normal")
    add_before(doc, intro, keywords, style="Normal")
    add_page_break_before(doc, intro)

    toc_heading = add_before(doc, intro, "TABLE OF CONTENTS", style="heading1")
    toc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_entries = [
        "1. Introduction",
        "2. Related Work",
        "3. Methodology",
        "4. Results and Discussion",
        "5. Conclusions and Recommendations",
        "References",
    ]
    for entry in toc_entries:
        add_before(doc, intro, entry, style="Normal")

    list_tables = add_before(doc, intro, "LIST OF TABLES", style="heading1")
    list_tables.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_captions = [p.text.strip() for p in doc.paragraphs if p.text.strip().startswith("Table ")]
    for caption in table_captions:
        add_before(doc, intro, caption, style="Normal")

    list_figures = add_before(doc, intro, "LIST OF FIGURES", style="heading1")
    list_figures.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_captions = [p.text.strip() for p in doc.paragraphs if p.text.strip().startswith("Fig. ")]
    for caption in figure_captions:
        add_before(doc, intro, caption, style="Normal")

    abbreviations_heading = add_before(doc, intro, "LIST OF ABBREVIATIONS", style="heading1")
    abbreviations_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    abbreviations = [
        ("AASIST", "Audio Anti-Spoofing using Integrated Spectro-Temporal Graph Attention Networks"),
        ("AWGN", "Additive White Gaussian Noise"),
        ("CI", "Confidence Interval"),
        ("CNN", "Convolutional Neural Network"),
        ("DET", "Detection Error Trade-off"),
        ("EER", "Equal Error Rate"),
        ("GRU", "Gated Recurrent Unit"),
        ("LSTM", "Long Short-Term Memory"),
        ("ONNX", "Open Neural Network Exchange"),
        ("PR-AUC", "Area Under the Precision-Recall Curve"),
        ("ROC-AUC", "Area Under the Receiver Operating Characteristic Curve"),
        ("RSS", "Resident Set Size"),
        ("RTF", "Real-Time Factor"),
        ("SHA-256", "Secure Hash Algorithm 256-bit"),
        ("STFT", "Short-Time Fourier Transform"),
    ]
    table = doc.add_table(rows=1, cols=2)
    if doc.tables and doc.tables[0].style is not None:
        table.style = doc.tables[0].style
    table.rows[0].cells[0].text = "Abbreviation"
    table.rows[0].cells[1].text = "Definition"
    for left, right in abbreviations:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
    intro._p.addprevious(table._tbl)
    add_page_break_before(doc, intro)

    # Anonymize document properties and force the required layout/typography.
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.subject = "BIT Genesis Research Awards 2026 submission"
    props.title = title

    # Remove the now-unused mailto relationship inherited from the conference
    # author block so anonymization holds at the package-metadata level too.
    for rel_id, relationship in list(doc.part.rels.items()):
        target = getattr(relationship, "target_ref", "") or ""
        if "hungngq@ueh.edu.vn" in target:
            del doc.part.rels[rel_id]

    configure_pages(doc)
    normalize_typography(doc)

    # Preserve a restrained hierarchy while respecting the 13 pt rule.
    for paragraph in doc.paragraphs:
        name = paragraph.style.name.lower() if paragraph.style else ""
        if name.startswith("heading"):
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.text.strip() in {"RESEARCH SUMMARY", "TABLE OF CONTENTS", "LIST OF TABLES", "LIST OF FIGURES", "LIST OF ABBREVIATIONS"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True

    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
