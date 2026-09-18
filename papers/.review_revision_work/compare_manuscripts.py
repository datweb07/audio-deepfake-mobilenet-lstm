from pathlib import Path
from zipfile import ZipFile
from collections import Counter
import re

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"D:\audio-deepfake-mobilenet-lstm\papers")
FILES = ["PaperID 804.docx", "PaperID 804 final.docx"]


def pt(value):
    return None if value is None else round(value.pt, 2)


def paragraph_size(paragraph):
    direct = [pt(run.font.size) for run in paragraph.runs if run.text.strip() and run.font.size]
    if direct:
        return Counter(direct).most_common(1)[0][0]
    style_size = paragraph.style.font.size if paragraph.style else None
    return pt(style_size)


def captions(doc):
    output = []
    for index, paragraph in enumerate(doc.paragraphs, 1):
        text = paragraph.text.strip()
        if re.match(r"^(Fig(?:ure)?\.?|Table)\s*\d+", text, re.I):
            output.append((index, paragraph.style.name if paragraph.style else "", paragraph_size(paragraph), text))
    return output


def image_paragraphs(doc):
    output = []
    for index, paragraph in enumerate(doc.paragraphs, 1):
        drawings = paragraph._p.xpath(".//w:drawing")
        if drawings:
            output.append((index, len(drawings), paragraph.text.strip()))
    return output


def linked_images(doc):
    output = []
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            for ci, cell in enumerate(row.cells, 1):
                count = len(cell._tc.xpath(".//w:drawing"))
                if count:
                    output.append((ti, ri, ci, count, cell.text.strip()[:100]))
    return output


for filename in FILES:
    path = ROOT / filename
    doc = Document(path)
    print(f"\n=== {filename} ===")
    print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "inline_shapes", len(doc.inline_shapes), "sections", len(doc.sections))
    print("table dims", [(len(t.rows), len(t.columns)) for t in doc.tables])
    print("image paragraphs", image_paragraphs(doc))
    print("images inside tables", linked_images(doc))
    print("captions")
    for row in captions(doc):
        print(row)
    print("heading/body size samples")
    seen = 0
    for index, paragraph in enumerate(doc.paragraphs, 1):
        text = paragraph.text.strip()
        if text and (paragraph.style.name.startswith("Heading") or seen < 18):
            print(index, paragraph.style.name, paragraph_size(paragraph), text[:150])
            seen += 1
            if seen >= 35:
                break
    with ZipFile(path) as zf:
        xml = zf.read("word/document.xml")
        print("OOXML drawing nodes", xml.count(b"<w:drawing"), "OMML", xml.count(b"<m:oMath"))

