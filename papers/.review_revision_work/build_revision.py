from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import shutil
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile

import numpy as np
import pandas as pd
from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"D:\audio-deepfake-mobilenet-lstm")
PAPERS = ROOT / "papers"
PAPER_IN = PAPERS / "PaperID 804.docx"
RESPONSE_TEMPLATE = PAPERS / "Response_to_Reviewers.docx"
PAPER_OUT = PAPERS / "PaperID 804 final.docx"
RESPONSE_OUT = PAPERS / "Response_to_Reviewers_final.docx"
CI_CSV = PAPERS / "tables" / "table_11_bootstrap_ci.csv"

SCORE_FILES = {
    "MobileNetV3": ROOT / "outputs/lava_5/clean/mobilenetv3_lstm/scores.csv",
    "ShuffleNetV2": ROOT / "outputs/lava_6/clean/shufflenetv2_lstm/scores.csv",
    "MnasNet-A1": ROOT / "outputs/lava_5/clean/mnasnet_lstm/scores.csv",
    "EffNet-B0 warm-up": ROOT / "outputs/lava_5/clean/efficientnet_b0_lstm/scores.csv",
    "RawNet2 ext.": ROOT / "outputs/lava_5/clean/rawnet2/scores.csv",
    "AASIST ext.": ROOT / "outputs/lava_5/clean/aasist/scores.csv",
}

MEDIA_REPLACEMENTS = {
    "word/media/image1.jpeg": PAPERS / "figures/lava_6_model_overview.png",
    "word/media/image2.jpeg": PAPERS / "figures/dataset_integrity_pipeline.png",
    "word/media/image3.jpeg": PAPERS / "figures/lightweight_temporal_pipeline.png",
    "word/media/image4.jpeg": PAPERS / "figures/training_provenance_strategies.png",
    "word/media/image5.jpeg": PAPERS / "figures/roc_comparison_6_models.png",
    "word/media/image6.jpeg": PAPERS / "figures/det_comparison_6_models.png",
    "word/media/image7.png": PAPERS / "figures/robustness_heatmap_6_models.png",
    "word/media/image8.jpeg": PAPERS / "figures/noise_f1_vs_snr_6_models.png",
    "word/media/image9.jpeg": PAPERS / "figures/pareto_eer_rtf_6_models.png",
}


def norm(text: str) -> str:
    return " ".join(text.split())


def find_paragraph(doc: Document, prefix: str):
    for p in doc.paragraphs:
        if norm(p.text).startswith(prefix):
            return p
    raise ValueError(f"Paragraph not found: {prefix}")


def replace_text(paragraph, text: str) -> None:
    paragraph.text = text


def new_paragraph_near(reference, text: str, style: str, before: bool = False):
    p = reference._parent.add_paragraph(text, style=style)
    p._p.getparent().remove(p._p)
    if before:
        reference._p.addprevious(p._p)
    else:
        reference._p.addnext(p._p)
    return p


def remove_paragraph(paragraph) -> None:
    paragraph._p.getparent().remove(paragraph._p)


def remove_table(table) -> None:
    table._tbl.getparent().remove(table._tbl)


def remove_image_paragraph(doc: Document, target_ref: str) -> None:
    for paragraph in list(doc.paragraphs):
        rel_ids = paragraph._p.xpath(".//a:blip/@r:embed")
        if any(doc.part.rels[rel_id].target_ref == target_ref for rel_id in rel_ids):
            remove_paragraph(paragraph)
            return
    raise ValueError(f"Image paragraph not found: {target_ref}")


def remove_equation_paragraph(doc: Document, token: str) -> None:
    for paragraph in list(doc.paragraphs):
        if token in "".join(paragraph._p.itertext()):
            remove_paragraph(paragraph)
            return
    raise ValueError(f"Equation paragraph not found: {token}")


def compact_table_numbering(doc: Document) -> None:
    mapping = {2: 1, 3: 2, 4: 3, 5: 4, 6: 5}
    for paragraph in doc.paragraphs:
        text = paragraph.text
        for old in mapping:
            text = text.replace(f"Table {old}", f"Table@@{old}@@")
        for old, new in mapping.items():
            text = text.replace(f"Table@@{old}@@", f"Table {new}")
        if text != paragraph.text:
            replace_text(paragraph, text)


def as_binary(series):
    if series.dtype == object:
        return series.map({"REAL": 0, "FAKE": 1}).astype(int)
    return series.astype(int)


def supplementary_metrics():
    result = {}
    for name, path in SCORE_FILES.items():
        frame = pd.read_csv(path)
        y = as_binary(frame["true_label"])
        pred = as_binary(frame["predicted_label"])
        score = frame["p_fake"].astype(float)
        tn = int(((y == 0) & (pred == 0)).sum())
        fp = int(((y == 0) & (pred == 1)).sum())
        y_values = y.to_numpy()
        score_values = score.to_numpy()
        order = np.argsort(-score_values, kind="mergesort")
        ranked_y = y_values[order]
        ranked_score = score_values[order]
        tp = np.cumsum(ranked_y == 1)
        fp_curve = np.cumsum(ranked_y == 0)
        threshold_ends = np.r_[np.where(np.diff(ranked_score))[0], ranked_y.size - 1]
        precision_curve = tp[threshold_ends] / (tp[threshold_ends] + fp_curve[threshold_ends])
        recall_curve = tp[threshold_ends] / max(1, int((y_values == 1).sum()))
        precision_curve = np.r_[1.0, precision_curve]
        recall_curve = np.r_[0.0, recall_curve]
        result[name] = {
            "specificity": tn / (tn + fp),
            "pr_auc": float(np.trapezoid(precision_curve, recall_curve)),
        }
    return result


def format_ci(lo, hi):
    return f"[{float(lo):.4f}, {float(hi):.4f}]"


def set_table_cell(cell, text: str, *, bold=False, size=7.0):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)


def add_ci_table(doc: Document, clean_table) -> None:
    ci = pd.read_csv(CI_CSV).set_index("Model")
    keys = ["MobileNetV3", "ShuffleNetV2", "MnasNet-A1", "EfficientNet-B0", "RawNet2", "AASIST"]
    caption = doc.add_paragraph(
        "Table 5. Full-test stratified percentile bootstrap 95% confidence intervals (1,000 iterations; seed 42). These are test-set intervals, not training-seed uncertainty.",
        style="Caption",
    )
    table = doc.add_table(rows=7, cols=4)
    table.style = doc.tables[3].style
    table.autofit = False
    for c, value in zip(table.rows[0].cells, ["Artifact", "F1 95% CI", "ROC-AUC 95% CI", "EER 95% CI"]):
        set_table_cell(c, value, bold=True, size=7.5)
    labels = {"MobileNetV3": "MobileNetV3", "ShuffleNetV2": "ShuffleNetV2", "MnasNet-A1": "MnasNet-A1", "EfficientNet-B0": "EffNet-B0 warm-up", "RawNet2": "RawNet2 ext.", "AASIST": "AASIST ext."}
    for row, key in zip(table.rows[1:], keys):
        vals = ci.loc[key]
        data = [labels[key], format_ci(vals.F1_lower, vals.F1_upper), format_ci(vals.AUC_lower, vals.AUC_upper), format_ci(vals.EER_lower, vals.EER_upper)]
        for c, value in zip(row.cells, data):
            set_table_cell(c, value, size=7.0)
    caption._p.getparent().remove(caption._p)
    table._tbl.getparent().remove(table._tbl)
    clean_table._tbl.addnext(caption._p)
    caption._p.addnext(table._tbl)


def replace_embedded_media(docx_path: Path) -> None:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, dir=docx_path.parent) as fh:
        tmp = Path(fh.name)
    try:
        with ZipFile(docx_path, "r") as zin, ZipFile(tmp, "w", ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename not in MEDIA_REPLACEMENTS:
                    zout.writestr(item, zin.read(item.filename))
                    continue
                image = Image.open(MEDIA_REPLACEMENTS[item.filename]).convert("RGB")
                buf = BytesIO()
                if item.filename.lower().endswith(".png"):
                    image.save(buf, format="PNG", dpi=(300, 300), optimize=True)
                else:
                    image.save(buf, format="JPEG", quality=95, dpi=(300, 300), optimize=True)
                zout.writestr(item, buf.getvalue())
        tmp.replace(docx_path)
    finally:
        if tmp.exists():
            tmp.unlink()


REFERENCES = [
    "[1] J. Yi et al., ‘Audio Deepfake Detection: A Survey,’ arXiv:2308.14970, 2023. https://doi.org/10.48550/arXiv.2308.14970.",
    "[2] M. Li, Y. Ahmadiadli, and X.-P. Zhang, ‘A Survey on Speech Deepfake Detection,’ ACM Computing Surveys, vol. 57, no. 7, pp. 1-38, 2025. https://doi.org/10.1145/3711898.",
    "[3] M. Todisco et al., ‘ASVspoof 2019: Future Horizons in Spoofed and Fake Audio Detection,’ Proc. Interspeech, pp. 1008-1012, 2019. https://doi.org/10.21437/Interspeech.2019-2249.",
    "[4] J. Frank and L. Schönherr, ‘WaveFake: A Data Set to Facilitate Audio Deepfake Detection,’ NeurIPS Datasets and Benchmarks, 2021, arXiv:2111.02813. https://doi.org/10.48550/arXiv.2111.02813.",
    "[5] T. Kinnunen et al., ‘t-DCF: A Detection Cost Function for the Tandem Assessment of Spoofing Countermeasures and Automatic Speaker Verification,’ Proc. Odyssey, pp. 312-319, 2018. https://doi.org/10.21437/Odyssey.2018-44.",
    "[6] J. Yamagishi et al., ‘ASVspoof 2021: Accelerating Progress in Spoofed and Deepfake Speech Detection,’ Proc. ASVspoof 2021 Workshop, pp. 47-54, 2021. https://doi.org/10.21437/ASVSPOOF.2021-8.",
    "[7] P. Veličković et al., ‘Graph Attention Networks,’ Proc. ICLR, 2018. https://openreview.net/forum?id=rJXMpikCZ.",
    "[8] H. Tak et al., ‘End-to-End Anti-Spoofing with RawNet2,’ Proc. ICASSP, pp. 6369-6373, 2021. https://doi.org/10.1109/ICASSP39728.2021.9414234.",
    "[9] M. Ravanelli and Y. Bengio, ‘Speaker Recognition from Raw Waveform with SincNet,’ Proc. IEEE SLT Workshop, pp. 1021-1028, 2018. https://doi.org/10.1109/SLT.2018.8639585.",
    "[10] J.-w. Jung et al., ‘AASIST: Audio Anti-Spoofing Using Integrated Spectro-Temporal Graph Attention Networks,’ Proc. ICASSP, pp. 6367-6371, 2022. https://doi.org/10.1109/ICASSP43922.2022.9747766.",
    "[11] H. Tak et al., ‘Automatic Speaker Verification Spoofing and Deepfake Detection Using wav2vec 2.0 and Data Augmentation,’ Proc. Odyssey, pp. 112-119, 2022. https://doi.org/10.21437/Odyssey.2022-16.",
    "[12] Z. Wu et al., ‘Light Convolutional Neural Network with Feature Genuinization for Detection of Synthetic Speech Attacks,’ Proc. Interspeech, pp. 1101-1105, 2020. https://doi.org/10.21437/Interspeech.2020-1810.",
    "[13] O. Pascu et al., ‘Towards Generalisable and Calibrated Audio Deepfake Detection with Self-Supervised Representations,’ Proc. Interspeech, pp. 4828-4832, 2024. https://doi.org/10.21437/Interspeech.2024-1302.",
    "[14] N. M. Müller, P. Czempin, F. Dieckmann, A. Froghyar, and K. Böttinger, ‘Does Audio Deepfake Detection Generalize?’ Proc. Interspeech, pp. 2783-2787, 2022. https://doi.org/10.21437/Interspeech.2022-108.",
    "[15] P. Kawa, M. Plata, and P. Syga, ‘Attack Agnostic Dataset: Towards Generalization and Stabilization of Audio Deepfake Detection,’ Proc. Interspeech, pp. 4023-4027, 2022. https://doi.org/10.21437/Interspeech.2022-10078.",
    "[16] N. M. Müller et al., ‘Harder or Different? Understanding Generalization of Audio Deepfake Detection,’ Proc. Interspeech, pp. 2705-2709, 2024. https://doi.org/10.21437/Interspeech.2024-247.",
    "[17] Y. Li et al., ‘Cross-Domain Audio Deepfake Detection: Dataset and Analysis,’ Proc. EMNLP, pp. 4977-4983, 2024. https://doi.org/10.18653/v1/2024.emnlp-main.286.",
    "[18] A. Howard et al., ‘Searching for MobileNetV3,’ Proc. ICCV, pp. 1314-1324, 2019. https://doi.org/10.1109/ICCV.2019.00140.",
    "[19] N. Ma et al., ‘ShuffleNet V2: Practical Guidelines for Efficient CNN Architecture Design,’ Proc. ECCV, pp. 116-131, 2018. https://doi.org/10.1007/978-3-030-01264-9_8.",
    "[20] M. Sandler et al., ‘MobileNetV2: Inverted Residuals and Linear Bottlenecks,’ Proc. CVPR, pp. 4510-4520, 2018. https://doi.org/10.1109/CVPR.2018.00474.",
    "[21] X. Zhang et al., ‘ShuffleNet: An Extremely Efficient Convolutional Neural Network for Mobile Devices,’ Proc. CVPR, pp. 6848-6856, 2018. https://doi.org/10.1109/CVPR.2018.00716.",
    "[22] M. Tan et al., ‘MnasNet: Platform-Aware Neural Architecture Search for Mobile,’ Proc. CVPR, pp. 2820-2828, 2019. https://doi.org/10.1109/CVPR.2019.00293.",
    "[23] M. Tan and Q. V. Le, ‘EfficientNet: Rethinking Model Scaling for Convolutional Neural Networks,’ Proc. ICML, PMLR, vol. 97, pp. 6105-6114, 2019. https://proceedings.mlr.press/v97/tan19a.html.",
    "[24] B. McFee et al., ‘librosa: Audio and Music Signal Analysis in Python,’ Proc. 14th Python in Science Conference, pp. 18-25, 2015. https://doi.org/10.25080/Majora-7b98e3ed-003.",
    "[25] A. Das et al., ‘Generalizable Audio Spoofing Detection Using Non-Semantic Representations,’ Proc. Interspeech, pp. 4553-4557, 2025. https://doi.org/10.21437/Interspeech.2025-1555.",
    "[26] N. Müller et al., ‘Replay Attacks Against Audio Deepfake Detection,’ Proc. Interspeech, pp. 2245-2249, 2025. https://doi.org/10.21437/Interspeech.2025-20.",
    "[27] I.-P. Ciobanu et al., ‘XMAD-Bench: Cross-Domain Multilingual Audio Deepfake Benchmark,’ Findings of EACL, pp. 3109-3120, 2026. https://doi.org/10.18653/v1/2026.findings-eacl.162.",
]


def revise_paper() -> None:
    shutil.copy2(PAPER_IN, PAPER_OUT)
    doc = Document(PAPER_OUT)
    extra = supplementary_metrics()

    abstract = find_paragraph(doc, "Abstract.")
    replace_text(abstract, "Abstract. LAVA compares heterogeneous deepfake-voice detectors under common integrity, P(FAKE), stress-diagnostic, and runtime contracts. Four locally trained Mel-sequence artifacts and externally pretrained RawNet2 and AASIST references were evaluated on a checksum-group-disjoint test set of 2,737 recordings. ShuffleNetV2 achieved F1 0.9824, ROC-AUC 0.9929, PR-AUC 0.9898, specificity 0.9848, and EER 1.46%. On one single-thread desktop CPU, MobileNetV3 had the lowest end-to-end latency (43.81 ms; RTF 0.0146). A fixed 100-recording diagnostic covered four AWGN levels, four codec round trips, and simulated replay; low-SNR noise was the main observed weakness. MobileNetV3 and ShuffleNetV2 were non-dominated under the evaluated objectives. The evidence supports offline deployment-oriented comparison, not full-test robustness, unseen-corpus, physical-replay, streaming, edge-device, or multi-seed claims.")

    contrib = find_paragraph(doc, "LAVA contributes:")
    replace_text(contrib, "The study makes three evidence-backed contributions:")
    anchor = contrib
    for text in [
        "• A checksum-group-disjoint, artifact-traceable clean evaluation contract centered on ROC-AUC/EER, calibrated score semantics, and explicit provenance.",
        "• Matched fixed-subset diagnostics for multiple noise levels, codec settings, and simulated replay, with aggregate failure-pattern analysis.",
        "• Single-thread desktop-CPU timing and exploratory three-objective Pareto analysis for scenario-specific artifact selection.",
    ]:
        anchor = new_paragraph_near(anchor, text, "Normal (Web)")

    intro = find_paragraph(doc, "Synthetic and converted speech pose risks")
    replace_text(intro, "Synthetic and converted speech pose risks to speaker verification and human decision-making. Surveys document rapid progress in deepfake voice detection [1], [2]. Shared resources such as ASVspoof and WaveFake support evaluation [3], [4]. Comparisons nevertheless remain difficult because studies differ in partitions, labels, duration, preprocessing, calibration, and runtime boundaries; byte-identical recordings across splits add an integrity risk.")
    related1 = find_paragraph(doc, "ASVspoof 2019 separates")
    replace_text(related1, "ASVspoof 2019 separates logical-access synthesis and conversion from replay-oriented physical access [3], [5]. ASVspoof 2021 adds channel, compression, and real-environment replay variation [6]. WaveFake covers multiple generators and two languages [4]. RawNet2 and SincNet represent raw-waveform modeling [8], [9]. AASIST and wav2vec 2.0 broaden graph and self-supervised representations [10], [11]. LCNN provides a complementary synthetic-speech baseline [12]. Scores remain tied to each dataset, attack partition, preprocessing, and decision protocol.")
    robustness = find_paragraph(doc, "Robustness and generalization concern different shifts.")
    replace_text(robustness, "Robustness and generalization concern different shifts. ASVspoof tasks cover simulated or real-environment replay and channel variation [3], [6]. A wav2vec 2.0 system addresses unmatched codec, attack, and source-domain variation [11]. Cross-source studies test generalization beyond one partition [13], [14]. Other work examines attack-agnostic data and the nature of domain shift [15], [16]. Cross-domain corpora extend this evidence [17]. Recent studies examine non-semantic representations and physical speaker-microphone replay [25], [26]. Multilingual cross-domain evaluation is also emerging [27]. These protocols are not interchangeable.")
    deployment = find_paragraph(doc, "Deployment evidence requires")
    replace_text(deployment, "Deployment evidence requires a common workload and timing boundary. MobileNet variants report efficiency for vision workloads [18], [20]. ShuffleNet variants do likewise [19], [21]. MnasNet and EfficientNet target mobile or compound scaling [22], [23]. These are not complete audio-pipeline measurements. Anti-spoofing studies mainly emphasize EER or t-DCF [8], [10]. Wav2vec-based work also focuses on detection rather than matched end-to-end timing [11]. Model size alone is not a deployment proxy.")
    gap = find_paragraph(doc, "The surveyed studies assess clean discrimination")
    replace_text(gap, "The surveyed studies assess clean discrimination, robustness or generalization, and efficiency under different protocols, making deployment-oriented comparison across heterogeneous artifacts difficult. LAVA addresses this comparability gap with one artifact-traceable offline contract; it does not claim unseen-corpus generalization, physical-replay robustness, streaming, or edge-device validation.")

    remove_table(doc.tables[0])
    remove_paragraph(find_paragraph(doc, "Table 1. Reported evaluation coverage"))
    remove_paragraph(find_paragraph(doc, "Note. Yes = evaluation results reported"))

    framework = find_paragraph(doc, "LAVA separates canonical data")
    replace_text(framework, "LAVA separates canonical data and integrity, registered detectors, conditions, measurements, and analyses. The registry records each artifact's input, duration, threshold, and provenance; adapters preserve native computation while exposing common scores and resource measures. Figure 1 summarizes the integrity-controlled evaluation flow.")

    dataset = find_paragraph(doc, "The internal collection lacks speaker")
    replace_text(dataset, "The repository contains an internally assembled collection, but its manifest does not retain source-dataset, speaker, source, generator, parent-recording, or dataset identifiers; those provenance fields are therefore reported as UNKNOWN rather than inferred. LAVA scanned 18,722 recordings (10,550 REAL; 8,172 FAKE). Files were grouped by exact byte-level SHA-256. Every member of a checksum group carrying both labels was quarantined: 14 such groups contained the 30 cross-label files. Within each remaining same-label group, the lexicographically first path was retained as the canonical representative and the other byte-identical copies were excluded. The resulting 18,232 recordings (10,493 REAL; 7,739 FAKE) were split with seed 42 into 12,762/2,733/2,737 train/validation/test files (Table 2). The supported property is checksum-group disjointness, not speaker-, source-, generator-, or corpus-disjointness. For h(x) = SHA256(x) and split group sets Gs,")
    remove_paragraph(find_paragraph(doc, "The manifest hash is:"))
    remove_paragraph(find_paragraph(doc, "8b55591d58d3658"))

    architecture = find_paragraph(doc, "For the lightweight family")
    replace_text(architecture, "The lightweight family uses TimeDistributed MobileNetV3Small, ShuffleNetV2-1.0x, MnasNet-A1-1.0, or EfficientNet-B0 segment encoders. Each encoder produces six embeddings, an LSTM with 128 units aggregates them, and a 64-unit ReLU layer plus dropout 0.4 feeds a sigmoid. RawNet2 instead uses a Sinc-style waveform front end, residual temporal blocks, attention, a GRU, and a two-class head [8], [9]. AASIST combines a raw-waveform front end with spectral-temporal graph attention, heterogeneous graph interaction, pooling, and two-class readout [7], [10]:")

    provenance = find_paragraph(doc, "MobileNet and EfficientNet use ImageNet")
    replace_text(provenance, "MobileNet and EfficientNet use ImageNet initialization; EfficientNet-B0 is the available warm-up checkpoint, not a completed fine-tuning result. MnasNet and ShuffleNet were trained from scratch, while RawNet2 and AASIST are external references (Table 2).")

    conversion = find_paragraph(doc, "RawNet2 and AASIST were strictly loaded")
    new_paragraph_near(conversion, "All local runs used batch size 16, seed 42, validation-loss checkpointing, early stopping, and ReduceLROnPlateau; Table 3 gives artifact-specific initialization, learning rates, regularization, stopping outcomes, and calibration. RawNet2 and AASIST were not locally trained; their external checkpoints were loaded in PyTorch and exported to ONNX.", "Normal", before=True)
    remove_paragraph(conversion)

    specification = doc.tables[1]
    specification_rows = [
        ["Detector", "Framework/input", "Architecture/params", "Initialization/provenance", "Training/artifact policy", "Threshold"],
        ["MobileNetV3-LSTM [18]", "TF/Keras; 6×Mel RGB, 3.0 s", "MobileNetV3Small → LSTM(128); 1.308M", "ImageNet; local", "50×10−4 warm-up; 50×10−5 last-20-layer fine-tune; BN frozen", ".82; val-F1"],
        ["ShuffleNetV2-LSTM [19]", "TF/Keras; 6×Mel RGB, 3.0 s", "ShuffleNetV2-1.0x → LSTM(128); 1.868M", "scratch; local", "Adam 3×10−4; end-to-end; 28 epochs, best 16", ".12; val-F1"],
        ["MnasNet-A1-LSTM [22]", "TF/Keras; 6×Mel RGB, 3.0 s", "MnasNet-A1 → LSTM(128); 3.369M", "scratch; local", "Adam 10−4; clip 1; WD 10−5; LS .1; 39 epochs, best 27", ".90; val-F1"],
        ["EffNet-B0 warm-up [23]", "TF/Keras; 6×Mel RGB, 3.0 s", "EfficientNet-B0 → LSTM(128); 4.779M", "ImageNet; local", "available warm-up checkpoint: epoch 47 at 10−4; fine-tune incomplete", ".90; val-F1"],
        ["RawNet2 [8]", "PyTorch→ONNX; waveform, 4.04 s", "Sinc + residual blocks + attention + GRU; 17.621M", "external reference", "no local training; exported checkpoint", ".50; default"],
        ["AASIST [10]", "PyTorch→ONNX; waveform, 4.04 s", "spectro-temporal graph attention; 0.298M", "external reference", "no local training; exported checkpoint", ".50; default"],
    ]
    for row_index, (row, values) in enumerate(zip(specification.rows, specification_rows)):
        for cell, value in zip(row.cells, values):
            set_table_cell(cell, value, bold=(row_index == 0), size=5.5 if row_index else 5.8)

    clean = doc.tables[2]
    # Add one editable column to the existing eight-column table.
    for row in clean.rows:
        row._tr.append(deepcopy(row.cells[-1]._tc))
    headers = ["Artifact", "Precision", "Recall", "Specificity", "F1 [95% CI]", "Macro-F1", "ROC-AUC [95% CI]", "PR-AUC", "EER [95% CI]"]
    for c, value in zip(clean.rows[0].cells, headers):
        set_table_cell(c, value, bold=True, size=6.3)
    existing = {
        "MobileNetV3": [".9741", ".9707", ".9724", ".9761", ".9911", ".0250"],
        "ShuffleNetV2": [".9795", ".9854", ".9824", ".9847", ".9929", ".0146"],
        "MnasNet-A1": [".9599", ".9690", ".9645", ".9690", ".9886", ".0310"],
        "EffNet-B0 warm-up": [".9724", ".9406", ".9563", ".9624", ".9877", ".0400"],
        "RawNet2 ext.": [".4380", ".6807", ".5330", ".4900", ".5178", ".4813"],
        "AASIST ext.": [".4758", ".5585", ".5139", ".5487", ".5597", ".4463"],
    }
    ci = pd.read_csv(CI_CSV).set_index("Model")
    ci_name = {"MobileNetV3": "MobileNetV3", "ShuffleNetV2": "ShuffleNetV2", "MnasNet-A1": "MnasNet-A1", "EffNet-B0 warm-up": "EfficientNet-B0", "RawNet2 ext.": "RawNet2", "AASIST ext.": "AASIST"}
    for row in clean.rows[1:]:
        label = norm(row.cells[0].text)
        precision, recall, f1, macro, roc, eer = existing[label]
        ci_row = ci.loc[ci_name[label]]
        f1_ci = f"{f1} [{ci_row['F1_lower']:.4f},{ci_row['F1_upper']:.4f}]"
        auc_ci = f"{roc} [{ci_row['AUC_lower']:.4f},{ci_row['AUC_upper']:.4f}]"
        eer_ci = f"{eer} [{ci_row['EER_lower']:.4f},{ci_row['EER_upper']:.4f}]"
        data = [label, precision, recall, f"{extra[label]['specificity']:.4f}".lstrip("0"), f1_ci, macro, auc_ci, f"{extra[label]['pr_auc']:.4f}".lstrip("0"), eer_ci]
        for c, value in zip(row.cells, data):
            set_table_cell(c, value, size=5.5)
    caption4 = find_paragraph(doc, "Table 4.")
    replace_text(caption4, "Table 4. Clean performance on all 2,737 canonical test recordings. Specificity, precision, recall, F1, and macro-F1 use each artifact's threshold; ROC-AUC, PR-AUC, and EER use raw P(FAKE) scores.")

    pareto = find_paragraph(doc, "The exploratory Pareto analysis minimizes")
    new_paragraph_near(pareto, "Uncertainty uses a stratified test-set percentile bootstrap (1,000 iterations; seed 42) for F1, ROC-AUC, and EER (Table 4). Paired correctness uses exact McNemar tests with Holm correction; this fixed-test evidence does not replace repeated training over independent seeds.", "Normal", before=True)

    results1 = find_paragraph(doc, "All artifacts passed load")
    replace_text(results1, "All artifacts passed load, score, hash, and adapter checks. ShuffleNetV2 had the highest ROC-AUC (0.9929), PR-AUC (0.9898), specificity (0.9848), and lowest EER (0.0146), followed by MobileNetV3, MnasNet-A1, and the available EfficientNet-B0 warm-up checkpoint (Table 4). Table 4 reports the corresponding bootstrap intervals.")
    results2 = find_paragraph(doc, "This is not a controlled architecture study.")
    replace_text(results2, "This is not a controlled architecture study. Training data, initialization, duration, representation, loaders, and calibration differ across artifacts; the results neither rank the RawNet2 or AASIST architectures nor estimate retrained performance. On fixed test predictions, ShuffleNetV2 differed from MobileNetV3 after Holm correction (adjusted p = 0.0096), and both differed strongly from the two external checkpoints; this is paired test-set evidence, not multi-seed training significance.")
    failure = find_paragraph(doc, "Aggregate failure-pattern analysis")
    replace_text(failure, "Because paired clean/stressed scores were not retained, prediction flips cannot be identified; the reported condition-level trends are aggregate failure patterns, not per-recording error analysis.")

    conclusion = find_paragraph(doc, "Conclusion")
    limitation = find_paragraph(doc, "Evidence is limited by heterogeneous training")
    conclusion._p.addprevious(limitation._p)
    limitations_heading = new_paragraph_near(limitation, "Limitations", "heading1", before=True)
    practical_heading = new_paragraph_near(limitations_heading, "Practical and Societal Implications", "heading1", before=True)
    practical = new_paragraph_near(practical_heading, "LAVA can support transparent artifact selection for defensive voice-authentication and media-verification screening, while retaining human review for consequential decisions. Its offline diagnostic evidence is not a guarantee of field security, demographic fairness, or edge-device readiness.", "p1a")
    limitations_heading._p.addprevious(practical._p)

    future = find_paragraph(doc, "LAVA provides an evidence-traceable comparison")
    replace_text(future, "LAVA compares six heterogeneous voice anti-spoofing artifacts under traceable contracts. MobileNetV3 and ShuffleNetV2 are non-dominated under the evaluated objectives, while low-SNR noise is the main weakness on the fixed diagnostic subset. Future work should add paired full-stress scores, multi-seed training, unseen corpora, physical replay, causal streaming, and representative edge hardware.")

    remove_image_paragraph(doc, "media/image1.jpeg")
    remove_image_paragraph(doc, "media/image3.jpeg")
    remove_image_paragraph(doc, "media/image4.jpeg")
    remove_image_paragraph(doc, "media/image5.jpeg")
    remove_image_paragraph(doc, "media/image6.jpeg")
    remove_image_paragraph(doc, "media/image8.jpeg")
    # The submitted caption contains a space before the period ("Fig. 1 .").
    remove_paragraph(find_paragraph(doc, "Fig. 1 ."))
    replace_text(find_paragraph(doc, "Fig. 2."), "Fig. 1. SHA-256 conflict quarantine, canonicalization, and checksum-group-disjoint splitting.")
    remove_paragraph(find_paragraph(doc, "(b) Artifact provenance"))
    remove_paragraph(find_paragraph(doc, "Fig. 3."))
    remove_paragraph(find_paragraph(doc, "ROC"))
    remove_paragraph(find_paragraph(doc, "DET/EER"))
    remove_paragraph(find_paragraph(doc, "Fig. 4."))
    replace_text(find_paragraph(doc, "Fig. 5."), "Fig. 2. Fixed-subset degradation across the nine diagnostic conditions; replay is simulated.")
    replace_text(find_paragraph(doc, "Fig. 6."), "Fig. 3. Exploratory Pareto space under the evaluated objectives; stars mark non-dominated artifacts.")
    for paragraph in doc.paragraphs:
        if "Figure 4" in paragraph.text:
            replace_text(paragraph, paragraph.text.replace("; Figure 4", "").replace("Figure 4", "Table 3"))
        elif "Figure 5" in paragraph.text:
            replace_text(paragraph, paragraph.text.replace("Figure 5", "Figure 2"))
        elif "Figure 6" in paragraph.text:
            replace_text(paragraph, paragraph.text.replace("Figure 6", "Figure 3"))
    remove_equation_paragraph(doc, "P=TPTP+FP")
    remove_equation_paragraph(doc, "FRR=FNFN+TP")
    remove_equation_paragraph(doc, "Xtm,k")
    remove_equation_paragraph(doc, "pfake=")
    remove_equation_paragraph(doc, "Q=NTtotal")
    compact_table_numbering(doc)

    ref_paras = [p for p in doc.paragraphs if p.text.strip().startswith("[")]
    ref_style = ref_paras[0].style.name
    for p in ref_paras:
        remove_paragraph(p)
    anchor = find_paragraph(doc, "References")
    for text in REFERENCES:
        anchor = new_paragraph_near(anchor, text, ref_style)
        anchor.paragraph_format.space_before = Pt(0)
        anchor.paragraph_format.space_after = Pt(0)
        anchor.paragraph_format.line_spacing = Pt(7.5)
        for run in anchor.runs:
            run.font.size = Pt(7.0)

    body = "\n".join(p.text for p in doc.paragraphs)
    assert len(doc.tables) == 5
    assert len(doc.inline_shapes) == 3
    assert len([p for p in doc.paragraphs if p.text.strip().startswith("[")]) == 27
    for token in ["2,737", "100-recording", "simulated replay", "UNKNOWN", "PR-AUC", "Specificity", "Limitations", "Practical and Societal Implications", "0.9929", "43.81"]:
        assert token in body, token
    doc.save(PAPER_OUT)
    replace_embedded_media(PAPER_OUT)


def add_response_item(doc: Document, number: str, comment: str, response: str, change: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run(f"Comment {number}. ").bold = True
    p.add_run(comment)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    p2.add_run("Response. ").bold = True
    p2.add_run(response)
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_after = Pt(6)
    p3.add_run("Changes in the manuscript. ").bold = True
    p3.add_run(change)


def clear_body_keep_sections(doc: Document):
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"):
            doc.element.body.remove(child)


def revise_response() -> None:
    shutil.copy2(RESPONSE_TEMPLATE, RESPONSE_OUT)
    doc = Document(RESPONSE_OUT)
    clear_body_keep_sections(doc)
    title = doc.add_paragraph("RESPONSE TO REVIEWERS", style="Heading 1")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("On behalf of all authors, we thank the reviewers for the careful and constructive evaluation. The manuscript has been revised point by point. Where a requested experiment is not supported by existing evidence, we strengthened the limitation and narrowed the claim instead of reporting unexecuted results. Location references use section/table/figure identifiers because pagination can vary slightly across Word versions.")

    doc.add_paragraph("Response to Reviewer 1", style="Heading 2")
    add_response_item(doc, "1", "Clearly specify the source datasets, recording counts, class balance, preprocessing pipeline, and criteria used to quarantine the 30 cross-label files.", "Addressed to the maximum supported by the repository. The manifest does not retain source-dataset identities, speakers, generators, or parent recordings; therefore, no dataset name was invented. We now state that the source is an internally assembled collection with provenance fields marked UNKNOWN, give scanned and retained class counts, explain every preprocessing parameter, and define the quarantine rule as exclusion of every file in an exact-SHA-256 group containing both labels.", "Section 3, Dataset Integrity and Canonical Split; Standardized Audio Preprocessing; Table 1.")
    add_response_item(doc, "2", "Report precision, recall, specificity, F1, ROC-AUC, PR-AUC, and confidence intervals in addition to AUC and EER.", "Addressed from stored full-test scores without re-running models. Table 3 now reports precision, recall, specificity, F1, macro-F1, ROC-AUC, PR-AUC, and EER for all six artifacts, with stratified percentile-bootstrap 95% confidence intervals for F1, ROC-AUC, and EER (1,000 iterations; seed 42). The text explicitly identifies these as test-set intervals rather than multi-seed uncertainty.", "Abstract; Sections 3 and 4; Table 3.")
    add_response_item(doc, "3", "Provide the exact architectures, pretrained weights, feature extraction settings, training hyperparameters, and calibration procedures for all six detector artifacts.", "Addressed. The revision specifies the four TimeDistributed backbones and shared LSTM head; RawNet2 and AASIST native waveform designs; exact Mel settings; local batch size, optimizer, learning rates, regularization and early-stopping outcomes; ImageNet versus scratch initialization; external-checkpoint provenance; ONNX parity; and validation-only versus default threshold sources. EfficientNet-B0 remains explicitly identified as the available warm-up checkpoint.", "Section 3, Standardized Audio Preprocessing, Evaluated Detector Architectures, and Training Artifact Provenance and Score Semantics; Table 2.")
    add_response_item(doc, "4", "Strengthen robustness evaluation by testing multiple SNR levels, codec bitrates, replay conditions, and unseen attack types rather than a fixed 100-recording diagnostic subset.", "Partially addressed, with scope kept truthful. The executed protocol already includes AWGN at 20/10/5/0 dB, MP3 at 128/64 kb/s, Opus at 64 kb/s, AAC at 96 kb/s, and a documented simulated-replay channel, all on the same stratified 100-recording subset. No unseen-attack or full-test stress outputs exist, so they were not fabricated. The manuscript now consistently calls this evidence diagnostic and lists unseen/full-test/physical-replay evaluation as future work.", "Abstract; Section 3, Clean and Diagnostic Stress Evaluation; Section 4, RQ2; Limitations; Conclusion.")
    add_response_item(doc, "5", "Perform multi-seed repeated experiments and statistical significance testing to establish whether differences are statistically reliable.", "Partially addressed without retraining. Stored paired full-test predictions support exact McNemar tests with Holm correction and stratified bootstrap intervals; these results are now summarized. Multi-seed training was not executed and cannot be inferred from one checkpoint per artifact. The manuscript therefore distinguishes paired test-set evidence from training-seed stability and retains multi-seed evaluation as a limitation and future-work item.", "Section 3, statistical protocol; Section 4, RQ1; Table 3; Limitations; Conclusion.")
    add_response_item(doc, "6", "Evaluate external-corpus and cross-dataset generalization, physical replay attacks, streaming performance, and resource consumption on representative edge devices before making deployment claims.", "Not executed; claims were narrowed rather than overstated. The repository has no external-corpus benchmark, physical speaker-microphone replay, causal-streaming measurement, or representative edge-device run. The revision explicitly limits its claim to deployment-oriented offline comparison on one desktop CPU and makes these experiments prerequisites for future deployment validation.", "Abstract; Related Work; Section 3, efficiency/stress protocols; Practical and Societal Implications; Limitations; Conclusion.")

    doc.add_paragraph("Response to Reviewer 2", style="Heading 2")
    add_response_item(doc, "General", "Maintain natural academic language; exclude nonsensical or unsupported data; complete references; and ensure captions and in-text figure/table references.", "Addressed through an editorial and evidence-consistency pass. Unsupported claims were removed or qualified. Bibliographic entries now consistently include venue, year, pages, volume/issue where applicable, and DOI; arXiv/OpenReview items use persistent identifiers where journal volume/issue or DOI does not apply. All tables remain native editable Word tables, all figures have captions, and their mentions were cross-checked in the text.", "Entire manuscript; References [1]-[27].")
    add_response_item(doc, "1", "Strengthen the abstract with empirical or quantitative results.", "Addressed. The abstract now includes full-test F1, ROC-AUC, PR-AUC, specificity, and EER for ShuffleNetV2; MobileNetV3 latency and RTF; the diagnostic subset size and stress coverage; and the exploratory non-dominated set, together with limitations.", "Abstract.")
    add_response_item(doc, "2", "Explicitly list the key contributions in bullet-point format.", "Addressed with three concise, evidence-backed bullets.", "Introduction.")
    add_response_item(doc, "3", "Incorporate recent relevant studies, especially from 2024, 2025, and 2026 onward.", "Addressed. The revision discusses 2024 cross-domain and calibration work, a 2025 non-semantic representation study, a 2025 physical-replay study, and a 2026 multilingual cross-domain benchmark. Each source is cited only at the claim it supports; the compact narrative replaces the earlier oversized literature table.", "Related Work; References [13], [16], [17], [25]-[27].")
    add_response_item(doc, "4", "Include a dedicated Limitations section before the conclusion.", "Addressed by moving and consolidating the limitations into a dedicated section immediately before the Conclusion.", "Limitations.")
    add_response_item(doc, "5", "Strictly adhere to the 10-12 page Springer limit, preferably 12 pages.", "Addressed structurally by preserving the A4 single-column Springer layout, using compact editable tables, removing redundant visuals and equations, and avoiding a large appendix. The final pagination was verified in Microsoft Word before delivery.", "Whole manuscript.")
    add_response_item(doc, "6", "Ensure all images are high resolution and publication-ready.", "Addressed. The three retained embedded figures use the corresponding repository-native 300-320 dpi sources; redundant figures were removed to meet the page limit without reducing the resolution of the scientific visuals that remain.", "Figures 1-3.")
    add_response_item(doc, "7", "Provide all tables in editable format.", "Verified. All five tables are native Word tables; no table is embedded as a raster image.", "Tables 1-5.")
    add_response_item(doc, "8", "Format equations with a standard math editor.", "Verified. Every retained equation remains a native editable Office Math (OMML) object; no equation was replaced by an image.", "Methodology equations.")
    add_response_item(doc, "9", "Include a dedicated bullet-point contribution list in the Introduction.", "Addressed; this overlaps Comment 2 and is implemented once to avoid repetition.", "Introduction.")
    add_response_item(doc, "10", "Add a section on practical implications and societal benefits.", "Addressed with a dedicated section explaining defensive use in voice authentication and media verification, the role of human oversight, and the limits of current offline diagnostic evidence.", "Practical and Societal Implications.")

    doc.add_paragraph("Final Remarks", style="Heading 3")
    doc.add_paragraph("The revision does not retrain any detector or invent unseen-corpus, physical-replay, streaming, edge-device, or multi-seed results. All added quantitative values are derived from saved full-test score files or existing benchmark artifacts, and all limitations remain explicit.")
    doc.save(RESPONSE_OUT)


if __name__ == "__main__":
    revise_paper()
    revise_response()
    print(PAPER_OUT)
    print(RESPONSE_OUT)
