from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path

from docx import Document

PAPERS = Path(__file__).resolve().parents[1]


def size_of(run, paragraph):
    if run.font.size:
        return round(run.font.size.pt, 2)
    if paragraph.style and paragraph.style.font.size:
        return round(paragraph.style.font.size.pt, 2)
    return None


def add(counter, paragraph):
    for run in paragraph.runs:
        if run.text.strip():
            counter[size_of(run, paragraph)] += len(run.text)


def inspect(path: Path):
    doc = Document(path)
    buckets = defaultdict(Counter)
    in_refs = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if text == "References":
            in_refs = True
            add(buckets["reference_heading"], p)
            continue
        if in_refs:
            add(buckets["references"], p)
            continue
        if not text:
            continue
        style = p.style.name if p.style else ""
        if text.startswith("LAVA:"):
            category = "title"
        elif text.startswith("Abstract."):
            category = "abstract"
        elif text.startswith("Keywords:"):
            category = "keywords"
        elif style.startswith("Heading"):
            category = style
        elif style == "Caption":
            category = "captions"
        elif style in {"Author", "Affiliation", "Email"}:
            category = style
        else:
            category = "body"
        add(buckets[category], p)

    table_counter = Counter()
    per_table = []
    for ti, table in enumerate(doc.tables, 1):
        c = Counter()
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    add(c, p)
                    add(table_counter, p)
        per_table.append((ti, len(table.rows), len(table.columns), c))

    print(f"\n=== {path.name} ===")
    for name, counter in buckets.items():
        print(name, dict(counter))
    print("all_tables", dict(table_counter))
    for row in per_table:
        print("table", row)


for filename in ["PaperID 804.docx", "PaperID 804 final.docx"]:
    inspect(PAPERS / filename)
