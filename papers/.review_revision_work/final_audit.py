from __future__ import annotations

import re
from collections import Counter
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "PaperID 804 final.docx"
BASE = ROOT / "PaperID 804.docx"
BACKUP = ROOT / "PaperID 804 final_before_last_review.docx"
RESPONSE = ROOT / "Response_to_Reviewers_final.docx"


def effective_size(run, paragraph):
    if run.font.size:
        return round(run.font.size.pt, 2)
    if paragraph.style and paragraph.style.font.size:
        return round(paragraph.style.font.size.pt, 2)
    return None


def doc_stats(path: Path):
    doc = Document(path)
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
        "equations": len(doc.element.body.findall(".//" + qn("m:oMath")))
        + len(doc.element.body.findall(".//" + qn("m:oMathPara"))),
    }


for path in (PAPER, RESPONSE):
    try:
        with ZipFile(path) as zf:
            bad = zf.testzip()
            print(f"ZIP {path.name}: PASS entries={len(zf.namelist())} bad={bad}")
    except BadZipFile as exc:
        print(f"ZIP {path.name}: FAIL {exc}")

print("accepted", doc_stats(BASE))
print("previous_working", doc_stats(BACKUP))
print("final", doc_stats(PAPER))

doc = Document(PAPER)
title = next(p for p in doc.paragraphs if p.text.strip().startswith("LAVA:"))
abstract = next(p for p in doc.paragraphs if p.text.strip().startswith("Abstract."))
keywords = next(p for p in doc.paragraphs if p.text.strip().startswith("Keywords:"))

for label, paragraph in (("title", title), ("abstract", abstract), ("keywords", keywords)):
    sizes = sorted({effective_size(r, paragraph) for r in paragraph.runs if r.text.strip()})
    print(label, sizes)

body_sizes = Counter()
body_examples = []
for p in doc.paragraphs:
    text = p.text.strip()
    if not text or p in (title, abstract, keywords):
        continue
    if p.style and p.style.name in {"Title", "Author", "Affiliation", "Email", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Caption", "Reference Item"}:
        continue
    if text in {"Acknowledgement", "References"}:
        continue
    for run in p.runs:
        if not run.text.strip():
            continue
        size = effective_size(run, p)
        body_sizes[size] += len(run.text)
        if size != 10.0 and len(body_examples) < 15:
            body_examples.append((p.style.name if p.style else None, size, text[:100]))
print("body size chars", body_sizes)
print("non-10 body examples", body_examples)

body_text = "\n".join(p.text for p in doc.paragraphs)

sentence_violations = []
for sentence in re.split(r"(?<=[.!?])\s+", body_text):
    ids = set()
    for group in re.findall(r"\[((?:\d{1,2})(?:\s*[-,]\s*\d{1,2})*)\]", sentence):
        for part in re.split(r"[,\s]+", group.strip()):
            if not part:
                continue
            if "-" in part:
                a, b = map(int, part.split("-", 1))
                ids.update(range(a, b + 1))
            elif part.isdigit():
                ids.add(int(part))
    if len(ids) > 2:
        sentence_violations.append((sorted(ids), sentence[:180]))
print("citation sentence violations", len(sentence_violations), sentence_violations[:5])

table_citation_violations = []
for table_index, table in enumerate(doc.tables, 1):
    for row_index, row in enumerate(table.rows, 1):
        for cell_index, cell in enumerate(row.cells, 1):
            ids = {int(x) for x in re.findall(r"\[(\d{1,2})\]", cell.text)}
            if len(ids) > 2:
                table_citation_violations.append((table_index, row_index, cell_index, sorted(ids)))
print("table-cell citation violations", table_citation_violations)

refs = [p.text.strip() for p in doc.paragraphs if re.match(r"^\[\d+\]", p.text.strip())]
years = [int(y) for ref in refs for y in re.findall(r"\b(20\d{2})\b", ref)]
recent = [ref for ref in refs if re.search(r"\b202[4-6]\b", ref)]
print("references", len(refs), "recent_2024_2026", len(recent))
print("recent ids", [re.match(r"^\[(\d+)\]", ref).group(1) for ref in recent])

response = Document(RESPONSE)
response_text = "\n".join(p.text for p in response.paragraphs)
required = [
    "Reviewer Comment 1:",
    "Response:",
    "Revision in Manuscript:",
    "References [1]-[24]",
    "Figures 1-6",
    "Partially addressed",
    "Retained as a limitation",
]
print("response markers", {token: token in response_text for token in required})

all_text = body_text + "\n" + "\n".join(
    cell.text for table in doc.tables for row in table.rows for cell in row.cells
)
paper_required = {
    "dataset counts": all(token in all_text for token in ["18,722", "18,232", "10,493", "7,739", "12,762", "2,733", "2,737"]),
    "cross-label rule": all(token in all_text for token in ["exact byte-level SHA-256", "14 such groups", "30 cross-label files", "quarantined"]),
    "preprocessing": all(token.lower() in all_text.lower() for token in ["22,050 Hz", "3.0 s", "six chronological", "Hann STFT", "hop 512", "128 HTK-style Mel", "224 × 224"]),
    "metrics": all(token in all_text for token in ["Precision", "Recall", "Specificity", "PR-AUC", "EER [95% CI]"]),
    "six artifacts": all(token in all_text for token in ["MobileNetV3", "ShuffleNetV2", "MnasNet", "EffNet-B0", "RawNet2", "AASIST"]),
    "robustness scope": all(token in all_text for token in ["100-recording", "20, 10, 5, 0 dB", "MP3", "Opus", "AAC", "simulated replay"]),
    "statistics scope": all(token in all_text for token in ["1,000 iterations", "McNemar", "Holm correction", "multi-seed"]),
    "deployment limits": all(token in all_text for token in ["unseen-corpus", "physical replay", "streaming", "edge hardware"]),
    "reviewer-2 sections": all(token in all_text for token in ["Limitations", "Practical and Societal Implications"]),
}
print("paper reviewer evidence", paper_required)

